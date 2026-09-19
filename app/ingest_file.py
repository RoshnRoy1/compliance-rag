import os
from langchain_community.document_loaders import PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
import chromadb
from chromadb.utils import embedding_functions
from config import CHROMA_PATH, COLLECTION_NAME, CHUNK_SIZE, CHUNK_OVERLAP, DOCS_FOLDER


def ingest_single_file(filepath):
    """Ingest one file into the existing vector database. Returns chunk count."""

    # Load based on file type
    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".pdf":
        loader = PyPDFLoader(filepath)
        pages = loader.load()
    elif ext in [".txt", ".md"]:
        with open(filepath, "r", encoding="utf-8") as f:
            text = f.read()
        from langchain.schema import Document
        pages = [Document(page_content=text, metadata={"source": filepath, "page": 0})]
    elif ext == ".docx":
        try:
            import docx
            doc = docx.Document(filepath)
            text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            from langchain.schema import Document
            pages = [Document(page_content=text, metadata={"source": filepath, "page": 0})]
        except ImportError:
            raise ValueError("python-docx not installed. Run: pip install python-docx")
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    if not pages:
        raise ValueError("No content found in file")

    # Chunk
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_documents(pages)

    if not chunks:
        raise ValueError("No chunks created from file")

    # Embed and store
    ef = embedding_functions.DefaultEmbeddingFunction()
    client = chromadb.PersistentClient(path=CHROMA_PATH)
    collection = client.get_collection(name=COLLECTION_NAME, embedding_function=ef)

    # Get current count for unique IDs
    current_count = collection.count()

    batch_size = 100
    for i in range(0, len(chunks), batch_size):
        batch = chunks[i:i + batch_size]
        collection.add(
            documents=[c.page_content for c in batch],
            metadatas=[c.metadata for c in batch],
            ids=[f"chunk_{current_count + i + j}" for j, c in enumerate(batch)],
        )

    return len(chunks)